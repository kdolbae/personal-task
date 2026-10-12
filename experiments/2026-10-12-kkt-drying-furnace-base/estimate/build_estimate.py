"""금강변압기 진공건조로 설치 공사 견적서 (openpyxl). 단가는 노란 칸에 입력."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "견적서_금강변압기_진공건조로_설치.xlsx")
F = "Arial"
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="2F3E52")
SEC = PatternFill("solid", fgColor="DDE5EF")
YEL = PatternFill("solid", fgColor="FFF2A8")
TOT = PatternFill("solid", fgColor="FCE9DD")
WON = '#,##0;(#,##0);"-"'
NUM = '#,##0.0;(#,##0.0);"-"'

def f(bold=False, color="000000", size=10):
    return Font(name=F, bold=bold, color=color, size=size)
BLUE, GREEN = "0000FF", "008000"

wb = Workbook()

# ------------------------------------------------------------------ 자재산출
ms = wb.active
ms.title = "자재산출"
ms["A1"] = "베이스 자재 산출 (사다리형 일체 베이스 v3, 10,000×5,000×H800)"
ms["A1"].font = f(True, size=13)
ms["A2"] = "파란 글씨 = 입력값(설계 가정). 단위중량은 KS D 3502 H형강 공칭값 기준, 발주 전 규격표와 대조 필요."
ms["A2"].font = f(color="555555", size=9)
hdr = ["번호", "품명", "규격", "수량(EA)", "개당 길이(m)", "총 길이(m)", "단위중량(kg/m)", "중량(kg)", "도장면적 계수(㎡/m)", "도장면적(㎡)", "비고"]
for c, h in enumerate(hdr, 1):
    x = ms.cell(4, c, h); x.font = f(True, "FFFFFF"); x.fill = HEAD; x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); x.border = BOX
# rung length per station = sum of 3 bays between rails at y=300,1400,3600,4700 minus flange 125 each
rows = [
    ("상부 사다리 레일", "H-250×125×6×9", 4, 10.0, 29.6, 0.99, "10m 연속, 밀착 시 슬라이딩 레일"),
    ("상부 사다리 가로대", "H-150×75×5×7", 18, "=(1100+2200+1100-3*125)/1000/3", 14.0, 0.59, "6개소×3칸, 레일 상면과 높이 일치"),
    ("하부 사다리 레일", "H-200×100×5.5×8", 8, 4.7, 21.3, 0.79, "통로부(W600)에서 분할"),
    ("하부 사다리 가로대", "H-150×75×5×7", 18, "=(1100+2200+1100-3*125)/1000/3", 14.0, 0.59, "6개소×3칸"),
    ("기둥", "H-125×125×6.5×9", 24, 0.35, 23.8, 0.73, "레일 교차점"),
]
r = 5
for i, (nm, spec, q, ln, kg, pa, note) in enumerate(rows, 1):
    ms.cell(r, 1, i); ms.cell(r, 2, nm); ms.cell(r, 3, spec)
    ms.cell(r, 4, q).font = f(color=BLUE)
    ms.cell(r, 5, ln).font = f(color=BLUE)
    ms.cell(r, 6, f"=D{r}*E{r}")
    ms.cell(r, 7, kg).font = f(color=BLUE)
    ms.cell(r, 8, f"=F{r}*G{r}")
    ms.cell(r, 9, pa).font = f(color=BLUE)
    ms.cell(r, 10, f"=F{r}*I{r}")
    ms.cell(r, 11, note)
    r += 1
# plates: weight by formula from dimensions
plates = [
    ("보강 스티프너", "PL-9 (120×230 2매 + 90×184 2매 / 기둥 1개소)", 24, "=(2*120*230*9+2*90*184*9)*7.85/1000000", "기둥 위치 레일 웹 내부"),
    ("푸시 반력 브래킷", "PL-25 300×400 2매 + PL-16 100×300 4매", 1, "=(2*300*400*25+4*100*300*16)*7.85/1000000", "탈착식, 밀착 후 제거 (예시)"),
]
for nm, spec, q, unitkg, note in plates:
    i += 1
    ms.cell(r, 1, i); ms.cell(r, 2, nm); ms.cell(r, 3, spec)
    ms.cell(r, 4, q).font = f(color=BLUE)
    ms.cell(r, 7, unitkg); ms.cell(r, 7).number_format = NUM
    ms.cell(r, 8, f"=D{r}*G{r}")
    ms.cell(r, 11, note + " (G열 = 1조당 kg)")
    r += 1
last = r - 1
ms.cell(r, 2, "소계").font = f(True)
ms.cell(r, 8, f"=SUM(H5:H{last})").font = f(True)
ms.cell(r, 10, f"=SUM(J5:J{last})").font = f(True)
sub = r
r += 1
ms.cell(r, 2, "손실률"); ms.cell(r, 3, 0.05).font = f(color=BLUE); ms.cell(r, 3).number_format = "0%"
loss = r
r += 1
ms.cell(r, 2, "발주 중량 (손실 포함)").font = f(True)
ms.cell(r, 8, f"=H{sub}*(1+C{loss})").font = f(True)
ms.cell(r, 9, "톤"); ms.cell(r, 10, f"=H{r}/1000"); ms.cell(r, 10).number_format = "0.00"
order = r
for rr in range(5, order + 1):
    for c in range(1, 12):
        x = ms.cell(rr, c); x.border = BOX
        if x.font.color is None or x.font.color.rgb in (None, "FF000000"):
            x.font = f(x.font.bold)
        if c in (5, 6, 7, 9): x.number_format = x.number_format if c == 7 and rr > 9 else NUM
        if c in (8, 10): x.number_format = NUM if c == 10 else WON
    if rr >= sub: 
        for c in range(1, 12): ms.cell(rr, c).fill = TOT
ms.cell(loss, 3).number_format = "0%"
ms.cell(order, 10).number_format = "0.00"
ms.cell(r + 2, 1, "도장면적 계수는 H형강 외주 길이(2×높이 + 4×플랜지폭 − 2×웨브두께) 기준 개략값이에요.").font = f(color="555555", size=9)
for col, w in zip("ABCDEFGHIJK", [6, 20, 40, 9, 11, 11, 13, 12, 14, 12, 34]):
    ms.column_dimensions[col].width = w
ms.row_dimensions[4].height = 30
MS_ORDER_KG = f"자재산출!$H${order}"
MS_PAINT = f"자재산출!$J${sub}"

# ------------------------------------------------------------------ 내역서
ds = wb.create_sheet("내역서")
ds["A1"] = "공사 내역서"
ds["A1"].font = f(True, size=13)
ds["A2"] = "노란 칸(단가)만 입력하면 금액이 계산돼요. 파란 글씨 수량은 가정값이라 확정 후 수정하세요. 초록 글씨는 자재산출 시트에서 가져온 값이에요."
ds["A2"].font = f(color="555555", size=9)
hdr = ["번호", "공종", "규격·내용", "단위", "수량", "재료비 단가", "노무비 단가", "경비 단가", "재료비", "노무비", "경비", "합계", "비고"]
for c, h in enumerate(hdr, 1):
    x = ds.cell(4, c, h); x.font = f(True, "FFFFFF"); x.fill = HEAD; x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); x.border = BOX

sections = [
 ("1", "가설·준비", [
   ("현장조사·실측", "지반, 동선, 간섭물, 하부 용접선 위치", "식", 1, "", ""),
   ("이동·설치 작업계획서", "인양계획(크레인 위치·줄걸이·하중배분), 위험성평가", "식", 1, "", "작업계획서 별도 제출"),
   ("구조검토", "베이스 레일·기둥·지반 지지력 검토", "식", 1, "", "개략검토상 스키드가 2줄에만 닿으면 휨 부족"),
   ("지반 평판재하시험", "아스팔트 지반 지지력 확인", "회", 1, "", "필요 시"),
 ]),
 ("2", "베이스 제작·설치", [
   ("강재 (H형강·강판)", "사다리형 일체 베이스, 손실 5% 포함", "kg", "=" + MS_ORDER_KG, "", "green"),
   ("가공·조립·용접", "절단, 개선, 조립, 필렛 용접", "kg", "=" + MS_ORDER_KG, "", "green"),
   ("방청 도장", "동력공구 소지조정 + 방청 프라이머 1회", "㎡", "=" + MS_PAINT, "", "green"),
   ("베이스 운반·설치", "현장 30t 갠트리 사용 가정, 레벨링 포함", "식", 1, "", ""),
   ("슬라이딩면 처리", "레일 상면 그리스 또는 PTFE·스테인리스 시트", "식", 1, "", "마찰계수 0.1 기준 푸시력 약 4.5tf"),
 ]),
 ("3", "1·2호기 회전·이동·안착", [
   ("대형 크레인 (100t급)", "상단 4점 인양, 회전", "일", 2, "", "장비 사양·일수는 인양계획 확정 후 수정"),
   ("보조 크레인 (60t급)", "하단 2점 지지, 회전 보조", "일", 2, "", ""),
   ("줄걸이·인양 작업", "줄걸이 반장, 기능공", "인", 8, "", ""),
   ("인양 치공구", "샤클, 슬링, 스프레더 빔", "식", 1, "", ""),
   ("도어 분리·재설치", "1호기 도어", "식", 1, "", "필요 시"),
   ("배관·노즐 보양", "노즐, 플랜지, 통기구", "식", 1, "", ""),
 ]),
 ("4", "밀착·하부 접합", [
   ("밀착 장비 임대", "유압잭 또는 윈치 (용량 25tf 이상)", "일", 2, "", "마찰계수 0.5 기준 22.5tf"),
   ("밀착·정렬 작업", "절단면 사상, 간격 확보, 정렬", "인", 6, "", ""),
   ("접합부 용접", "하부 접합선 용접 (통로부 순높이 550)", "m", 10, "", "용접 길이 미확정, 10m는 가정"),
   ("비파괴 검사", "PT 또는 UT", "식", 1, "", "요구 시"),
   ("용접부 방청·도장 보수", "사상 후 프라이머·상도", "식", 1, "", ""),
 ]),
 ("5", "안전·부대", [
   ("안전관리", "신호수, 화기감시자, 출입통제", "인", 6, "", ""),
   ("밀폐·협소공간 작업 대책", "환기, 조명, 가스측정", "식", 1, "", "통로부 저자세 용접"),
   ("폐기물 처리·정리", "사상분진, 폐도료", "식", 1, "", ""),
 ]),
]
r = 5
item_rows = []
sec_rows = []
for sno, sname, items in sections:
    ds.cell(r, 1, sno); ds.cell(r, 2, sname)
    for c in range(1, 14):
        ds.cell(r, c).fill = SEC; ds.cell(r, c).font = f(True); ds.cell(r, c).border = BOX
    srow = r; r += 1
    first = r
    for k, (nm, spec, unit, qty, _, note) in enumerate(items, 1):
        ds.cell(r, 1, f"{sno}.{k}"); ds.cell(r, 2, nm); ds.cell(r, 3, spec); ds.cell(r, 4, unit)
        q = ds.cell(r, 5, qty)
        q.font = f(color=GREEN if note == "green" else BLUE)
        q.number_format = NUM if unit in ("kg", "㎡") else '#,##0.0;(#,##0.0);"-"' if unit == "m" else "#,##0"
        for c in (6, 7, 8):
            x = ds.cell(r, c); x.fill = YEL; x.number_format = WON; x.font = f(color=BLUE)
        ds.cell(r, 9, f"=E{r}*F{r}"); ds.cell(r, 10, f"=E{r}*G{r}"); ds.cell(r, 11, f"=E{r}*H{r}")
        ds.cell(r, 12, f"=SUM(I{r}:K{r})")
        for c in (9, 10, 11, 12): ds.cell(r, c).number_format = WON; ds.cell(r, c).font = f()
        ds.cell(r, 13, "자재산출 시트 연동" if note == "green" else note)
        for c in (1, 2, 3, 4, 13): ds.cell(r, c).font = f()
        for c in range(1, 14): ds.cell(r, c).border = BOX
        r += 1
    last = r - 1
    for c, col in ((9, "I"), (10, "J"), (11, "K"), (12, "L")):
        x = ds.cell(srow, c, f"=SUM({col}{first}:{col}{last})"); x.number_format = WON
    sec_rows.append(srow)
dsum = r
ds.cell(r, 2, "직접공사비 계")
for c, col in ((9, "I"), (10, "J"), (11, "K"), (12, "L")):
    ds.cell(r, c, "=" + "+".join(f"{col}{s}" for s in sec_rows)).number_format = WON
for c in range(1, 14):
    ds.cell(r, c).fill = TOT; ds.cell(r, c).font = f(True); ds.cell(r, c).border = BOX
for col, w in zip("ABCDEFGHIJKLM", [6, 22, 40, 6, 10, 12, 12, 12, 13, 13, 13, 14, 36]):
    ds.column_dimensions[col].width = w
ds.row_dimensions[4].height = 30
ds.freeze_panes = "A5"
SEC_NAMES = [(s[0], s[1]) for s in sections]

# ------------------------------------------------------------------ 견적서 (표지)
cs = wb.create_sheet("견적서", 0)
cs.merge_cells("A1:F1")
cs["A1"] = "견   적   서"; cs["A1"].font = f(True, size=20); cs["A1"].alignment = Alignment(horizontal="center")
info = [
    ("수신", "금강변압기 귀중"),
    ("공사명", "진공건조로 2기 회전·이동·안착 및 하부 접합 공사"),
    ("공사 장소", "금강변압기 공장 야적장"),
    ("견적일", "2026-10-12"),
    ("유효기간", "견적일로부터 30일"),
    ("공급자", ""),
]
r = 3
for k, v in info:
    cs.cell(r, 1, k).font = f(True); cs.cell(r, 2, v).font = f(color=BLUE if k in ("공급자",) else "000000")
    if k == "공급자": cs.cell(r, 2).fill = YEL
    cs.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    r += 1
r += 1
hdr = ["번호", "공종", "재료비", "노무비", "경비", "금액"]
for c, h in enumerate(hdr, 1):
    x = cs.cell(r, c, h); x.font = f(True, "FFFFFF"); x.fill = HEAD; x.alignment = Alignment(horizontal="center"); x.border = BOX
r += 1
first = r
for (sno, sname), srow in zip(SEC_NAMES, sec_rows):
    cs.cell(r, 1, sno); cs.cell(r, 2, sname)
    for c, col in ((3, "I"), (4, "J"), (5, "K"), (6, "L")):
        x = cs.cell(r, c, f"=내역서!{col}{srow}"); x.font = f(color=GREEN); x.number_format = WON
    for c in range(1, 7): cs.cell(r, c).border = BOX
    cs.cell(r, 1).font = f(); cs.cell(r, 2).font = f()
    r += 1
last = r - 1
def line(label, formula, bold=False, fill=None, rate=None):
    global r
    cs.cell(r, 2, label).font = f(bold)
    if rate is not None:
        x = cs.cell(r, 5, rate); x.font = f(color=BLUE); x.number_format = "0.0%"; x.fill = YEL
    x = cs.cell(r, 6, formula); x.number_format = WON; x.font = f(bold)
    for c in range(1, 7):
        cs.cell(r, c).border = BOX
        if fill: cs.cell(r, c).fill = fill
    r += 1
    return r - 1
direct = line("직접공사비 계", f"=SUM(F{first}:F{last})", True)
ga = line("일반관리비 (직접공사비 × 요율)", f"=F{direct}*E{r}", rate=0.06)
pf = line("이윤 ((직접공사비 + 일반관리비) × 요율)", f"=(F{direct}+F{ga})*E{r}", rate=0.10)
supply = line("공급가액", f"=F{direct}+F{ga}+F{pf}", True)
vat = line("부가가치세 (10%)", f"=F{supply}*0.1")
total = line("총 견적금액 (VAT 포함)", f"=F{supply}+F{vat}", True, TOT)
cs.cell(ga, 1).comment = Comment("일반관리비 6%, 이윤 10%는 예시 요율이에요. 회사 기준으로 바꾸세요.", "Claude")
r += 1
notes = [
    "특기사항",
    "1. 크레인 사양·대수·일수는 인양계획서와 구조검토 확정 후 조정합니다. 현장 30t 갠트리로는 1호기(45t) 단독 인양이 불가합니다.",
    "2. 베이스 강재는 개략 산출(손실 5% 포함 약 3.2톤)이며, 구조검토 결과에 따라 레일 줄 수·보강이 늘 수 있습니다.",
    "3. 하부 접합 용접 길이와 비파괴 검사 범위는 발주처 확인 후 정산합니다.",
    "4. 지반 지지력 부족 시 보강 비용은 별도입니다.",
    "5. 진공건조로 본체의 내부 작업, 전기·배관 해체·복구는 제외합니다.",
]
for i, t in enumerate(notes):
    cs.cell(r, 1, t).font = f(i == 0, size=10 if i else 11)
    cs.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    cs.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    cs.row_dimensions[r].height = 28 if i else 18
    r += 1
r += 1
cs.cell(r, 1, "입력 안내: 노란 칸 = 직접 입력(단가, 요율, 공급자). 파란 글씨 = 가정값. 초록 글씨 = 다른 시트에서 연동.").font = f(color="555555", size=9)
cs.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
for col, w in zip("ABCDEF", [10, 44, 15, 15, 15, 17]):
    cs.column_dimensions[col].width = w

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
cs.page_setup.orientation = "portrait"; cs.page_setup.fitToHeight = 1
wb.save(OUT)
print(OUT)
