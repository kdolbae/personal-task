import sys, subprocess, re, os, shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from pypdf import PdfReader, PdfWriter
FN="Pretendard"; FN_SB="Pretendard SemiBold"; QNO="NM-2610-01"
C_NAVY,C_TEXT,C_SUB,C_LINE,C_HEAD="0F2A44","1D2430","5D6878","DDE3EA","F4F6F9"
SUP_LINE="엔엠솔루션  |  경기도 평택시 고덕동 도시지원로 121 지식공장아이타워 906호  |  대표전화 1800-5901  |  직통 010-2369-3691  |  nm-solution.co.kr"
SRC=sys.argv[1]; OUT=sys.argv[2]; W='/tmp/claude-0/-home-claude-personal-task/cc458d43-e61f-5be7-bef2-0fa5339e22e4/scratchpad/pdfwork'
shutil.rmtree(W,ignore_errors=True); os.makedirs(W)
DETAIL={"A":8,"B":16,"C":22,"D":7,"E":7,"F":11,"G":13,"H":22,"I":7}
PRINT_W={"견적서":{"A":9,"B":13,"C":13,"D":22,"E":7,"F":6,"G":12,"H":15,"I":24},"AB비교":{"A":14,"B":20,"C":20,"D":15,"E":44},
         "산출근거":{"A":6,"B":13,"C":20,"D":30,"E":12,"F":30,"G":14},
         "도비":DETAIL,"용접":DETAIL,"베이스":DETAIL,"B_도비":DETAIL,"B_용접":DETAIL,"B_조방":DETAIL,
         "공정표":{"A":24,"B":26,"C":5.5,**{c:7 for c in "DEFGHIJKLMN"}},
         "요율근거":{"A":18,"B":22,"C":12,"D":42,"E":8}}
WRAP={"견적서":("B","D","I"),"AB비교":("B","C","E"),"산출근거":("B","C","D","F","G"),"요율근거":("A","B","C","D"),"공정표":("A","B"),
      **{k:("B","C","H") for k in ("도비","용접","베이스","B_도비","B_용접","B_조방")}}
AUTOROW_FROM={"견적서":17,"산출근거":2,"요율근거":2,"AB비교":4,"공정표":4,**{k:8 for k in ("도비","용접","베이스","B_도비","B_용접","B_조방")}}
def layout(ws):
    # 인쇄용 레이아웃: 세로 A4 폭에 맞게 열 너비를 줄이고, 긴 글은 줄바꿈 + 행 높이 자동
    from copy import copy
    w=PRINT_W.get(ws.title)
    if not w: return
    for col,v in w.items(): ws.column_dimensions[col].width=v
    cols=WRAP.get(ws.title,())
    r0=AUTOROW_FROM.get(ws.title,1)
    for row in ws.iter_rows(min_row=r0):
        for x in row:
            if getattr(x,"column_letter",None) in cols and isinstance(x.value,str) and len(x.value)>8:
                if ws.title=="요율근거" and x.column_letter=="A" and not any(x.coordinate in mr for mr in ws.merged_cells.ranges): continue
                a=copy(x.alignment); a.wrap_text=True; a.vertical='center'; x.alignment=a
    import math
    def textw(t): return sum(2.0 if ord(ch)>127 else 1.05 for ch in t)
    for r in range(r0, ws.max_row+1):
        need=0
        for x in ws[r]:
            cl=getattr(x,"column_letter",None)
            if cl in cols and isinstance(x.value,str) and x.value:
                if ws.title=="요율근거" and cl=="A" and not any(x.coordinate in mr for mr in ws.merged_cells.ranges): continue
                cw=w.get(cl, ws.column_dimensions[cl].width or 8.43)
                # 병합 셀이면 병합 범위 너비 합산
                for mr in ws.merged_cells.ranges:
                    if x.coordinate in mr and mr.min_row==mr.max_row:
                        from openpyxl.utils import get_column_letter as GL
                        cw=sum(w.get(GL(c), ws.column_dimensions[GL(c)].width or 8.43) for c in range(mr.min_col, mr.max_col+1))
                sz=(x.font.sz or 10)
                lines=sum(max(1, math.ceil(textw(seg)*sz/10/max(cw-1,1))) for seg in x.value.split("\n"))
                need=max(need, lines*(sz*1.28)+(6 if ws.title=="요율근거" else 3))
        rd=ws.row_dimensions[r]
        cur=rd.height
        if need and (cur is None or need>cur): rd.height=need
    if ws.title=="요율근거":
        from openpyxl.worksheet.pagebreak import Break
        for r in range(1, ws.max_row+1):
            v=ws.cell(r,1).value
            if isinstance(v,str) and v.startswith("시중노임"):
                ws.row_breaks.append(Break(id=r-1)); break
def prep(wb, index=None):
    names=[ws.title for ws in wb.worksheets]
    for i,ws in enumerate(wb.worksheets):
        layout(ws)
        ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.orientation="portrait"
        ws.print_options.verticalCentered=False
        ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_options.horizontalCentered=True
        ws.page_margins.left=ws.page_margins.right=0.63; ws.page_margins.top=ws.page_margins.bottom=0.71
        ws.page_margins.header=ws.page_margins.footer=0.35
        ws.oddHeader.left.text=f"&\"{FN}\"&8 금강변압기 진공건조로 회전·안착·용접 견적서  ·  {QNO}"
        ws.oddHeader.right.text=f"&\"{FN}\"&8 §{i+1}. &A"
        ws.oddFooter.left.text=f"&\"{FN}\"&8 {SUP_LINE}"
        ws.oddFooter.center.text=""
        ws.oddFooter.right.text=f"&\"{FN}\"&8 &P / &N"
    if index is not None:
        ix=wb.create_sheet("목차",0)
        ix.page_setup.paperSize=ix.PAPERSIZE_A4; ix.page_setup.fitToWidth=1; ix.page_setup.fitToHeight=0
        ix.sheet_properties.pageSetUpPr.fitToPage=True
        ix.page_margins.left=ix.page_margins.right=0.63; ix.page_margins.top=ix.page_margins.bottom=0.71
        ix.oddFooter.left.text=f"&\"{FN}\"&8 {SUP_LINE}"; ix.oddFooter.right.text=f"&\"{FN}\"&8 &P / &N"
        ix.column_dimensions['A'].width=8; ix.column_dimensions['B'].width=34; ix.column_dimensions['C'].width=46; ix.column_dimensions['D'].width=12
        ix['A1']="견적서 목차"; ix['A1'].font=Font(name=FN,bold=True,size=16,color=C_NAVY)
        ix['A2']="금강변압기 진공건조로 회전·안착 및 접합 용접 공사 (공급자: 엔엠솔루션)"; ix['A2'].font=Font(name=FN,size=10,color=C_SUB)
        hdr=["No","시트","내용","페이지"]
        desc={"견적서":"표지·공종 합계·특기사항","원가계산서":"비목별 원가 집계 (A안)","AB비교":"A안·B안 금액 비교","산출근거":"항목별 산출식·근거",
              "도비":"A안 도비 내역","용접":"A안 용접 내역","베이스":"A안 베이스 내역","B_도비":"B안 도비 내역","B_용접":"B안 용접 내역","B_조방":"B안 조방 개조 내역",
              "베이스자재":"H빔 자재 산출","공정표":"공정표 (가정 일정 10/21~11/2)","요율근거":"보험·관리비·이윤 요율 근거"}
        thin=Side(style='thin',color=C_LINE)
        for c,h in enumerate(hdr,1):
            x=ix.cell(4,c,h); x.font=Font(name=FN_SB,size=10,color=C_TEXT); x.fill=PatternFill('solid',fgColor=C_HEAD); x.alignment=Alignment(horizontal='center')
            x.border=Border(top=Side(style='medium',color=C_NAVY),bottom=thin,left=thin,right=thin)
        r=5
        for i,n in enumerate(names):
            vals=[f"§{i+1}",n,desc.get(n,""),f"p.{index[n][0]}" if index[n][0]==index[n][1] else f"p.{index[n][0]}~{index[n][1]}"]
            for c,v in enumerate(vals,1):
                x=ix.cell(r,c,v); x.font=Font(name=FN,size=10,color=C_TEXT); x.border=Border(top=thin,bottom=thin,left=thin,right=thin)
                x.alignment=Alignment(horizontal='center' if c in (1,4) else 'left',vertical='center')
            ix.row_dimensions[r].height=20; r+=1
    return names
def convert(xlsx):
    subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",W,xlsx],check=True,capture_output=True,timeout=180)
    return os.path.join(W,os.path.splitext(os.path.basename(xlsx))[0]+".pdf")
# pass 1
wb=load_workbook(SRC); names=prep(wb); p1=os.path.join(W,"pass1.xlsx"); wb.save(p1); pdf1=convert(p1)
rd=PdfReader(pdf1); pages={}
for pi,pg in enumerate(rd.pages,1):
    t=pg.extract_text() or ""
    m=re.search(r"§(\d+)\.",t)
    if m:
        n=names[int(m.group(1))-1]; pages.setdefault(n,[pi,pi]); pages[n][1]=pi
    else:
        print("no header on page",pi,repr(t[:60]))
print(pages)
# pass 2 with index (index page adds 1 page; assume index is 1 page)
index={n:[a+1,b+1] for n,(a,b) in pages.items()}
wb=load_workbook(SRC); prep(wb,index); p2=os.path.join(W,"pass2.xlsx"); wb.save(p2); pdf2=convert(p2)
rd=PdfReader(pdf2); w=PdfWriter()
for pg in rd.pages: w.add_page(pg)
w.add_outline_item("목차",0)
for i,n in enumerate(names):
    w.add_outline_item(f"§{i+1}. {n}",index[n][0]-1)
with open(OUT,"wb") as f: w.write(f)
print("pages",len(rd.pages),"->",OUT)
