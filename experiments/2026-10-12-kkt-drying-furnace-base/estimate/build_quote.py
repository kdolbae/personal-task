"""엔엠솔루션 → 금강변압기 견적서 (서해테크 도비 견적서 양식 참고).
공종 목표금액: 도비 2,000만 / 용접 2,500만 / 베이스 1,800만 (부가세 별도, 사용자 지정).
각 공종의 '기업이윤'은 목표금액에 맞추는 조정 항목이다."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "견적서_금강변압기_진공건조로_회전안착_용접_v2.xlsx")
FN = "맑은 고딕"
WON = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'
thin = Side(style="thin", color="000000")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GRAY = PatternFill("solid", fgColor="F2F2F2")
YEL = PatternFill("solid", fgColor="FFF2CC")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
INS = 0.033            # 고용·산재보험 요율 (서해테크 견적서와 같은 3.3%)
STEEL_KG = 3226.57     # 자재산출 시트 발주중량(손실 5% 포함), v3 베이스

def font(b=False, sz=11, color="000000"):
    return Font(name=FN, bold=b, size=sz, color=color)

# (구분, 항목, 규격, 단위, 수량, 단가, 비고, 인건비여부)
DOBI = [
    ("기계작업", "인건비", "도비 전문기공", "M", 21, 300000, "7명 × 3일 (안착 1일 + 밀착 이동 2일)", True),
    ("기계작업", "인건비", "안전관리자", "M", 3, 350000, "3일", True),
    ("기계작업", "크 레 인", "100TON", "대", 1, 2400000, "베이스 안착일 (1일), 일 240만 원", False),
    ("기계작업", "크 레 인", "60TON", "대", 1, 1700000, "베이스 안착일 (1일), 일 170만 원", False),
    ("기계작업", "밀착 장비", "유압잭·윈치·스키드", "SET", 1, 2000000, "밀착 이동 2일", False),
    ("기계작업", "공구사용손료", "도비전문장비", "SET", 1, 2500000, "", False),
    ("기계작업", "일반관리비", "", "DAY", 3, 700000, "", False),
    ("기타비용", "소모품잡자재", "", "식", 1, 1000000, "", False),
]
WELD = [
    ("용접작업", "인건비", "용접사", "M", 9, 350000, "3명 × 3일", True),
    ("용접작업", "인건비", "고급기술자", "M", 6, 450000, "2명 × 3일", True),
    ("사상작업", "인건비", "사상 기술자", "M", 4, 300000, "2명 × 2일", True),
    ("사상작업", "인건비", "조공", "M", 4, 200000, "2명 × 2일", True),
    ("용접작업", "인건비", "안전관리자·화기감시", "M", 5, 300000, "사상 2일 + 용접 3일", True),
    ("용접작업", "장비 대여", "CO2 용접기", "대·일", 9, 400000, "3대 × 3일, 대당 일 40만 원", False),
    ("용접작업", "장비사용료", "그라인더·환기·조명", "SET", 1, 800000, "", False),
    ("용접작업", "재료비", "CO2 가스·와이어·디스크", "식", 1, 3000000, "", False),
    ("용접작업", "비파괴검사", "PT/UT", "식", 1, 2500000, "요구 시, 범위 확인 필요", False),
    ("용접작업", "방청·도장 보수", "용접부", "식", 1, 1500000, "", False),
    ("기타비용", "일반관리비", "", "DAY", 5, 600000, "", False),
]
BASE = [
    ("자재", "H빔 구매 (중고)", "H-250×250, 10m 4EA·5m 16EA·0.6m 30EA", "식", 1, 9577783, "공급처 견적, 신품 시 13,575,440원", False),
    ("자재", "운임", "H빔 운반", "식", 1, 160000, "공급처 견적", False),
    ("자재", "지 게 차", "자재 하역·조립", "일", 2, 300000, "", False),
    ("제작작업", "인건비", "제작 기공", "M", 12, 300000, "4명 × 3일", True),
    ("제작작업", "장비사용료", "용접기·절단기", "SET", 1, 1000000, "", False),
    ("제작작업", "소모품", "용접봉·가스·디스크", "식", 1, 1000000, "", False),
    ("제작작업", "슬라이딩면 처리", "그리스·PTFE 시트", "식", 1, 500000, "밀착 이동용 레일 상면", False),
    ("기타비용", "일반관리비", "", "DAY", 3, 300000, "", False),
]
TARGETS = {"도비": 20_000_000, "용접": 25_000_000, "베이스": 18_000_000}

wb = Workbook()
cover = wb.active
cover.title = "견적서"

def header(ws, title_work, sub=None):
    ws.merge_cells("A1:I2"); ws["A1"] = "견      적      서"
    ws["A1"].font = font(True, 23); ws["A1"].alignment = CEN
    ws.merge_cells("A4:C4"); ws["A4"] = " 금강변압기 유한회사 귀하"; ws["A4"].font = font(True, 15)
    ws.merge_cells("A5:C5"); ws["A5"] = "   담당자 :"; ws["A5"].font = font()
    ws.merge_cells("G4:I4"); ws["G4"] = "엔엠솔루션"; ws["G4"].font = font(True, 15); ws["G4"].alignment = CEN
    ws.merge_cells("G5:I5"); ws["G5"] = "대표 / 연락처 :"; ws["G5"].font = font(); ws["G5"].fill = YEL
    ws["A6"] = "   견적일 : 2026-10-12"; ws["A6"].font = font()
    ws.merge_cells("A7:I7"); ws["A7"] = "   공사 장소 : 경기 화성시 마도면 마도공단로1길 214 금강변압기 유한회사(KUMKANG TRANSFORMER) 마당"; ws["A7"].font = font()
    ws.merge_cells("A8:I8"); ws["A8"] = "  작업 내용: " + title_work; ws["A8"].font = font(True, 12); ws["A8"].alignment = LEFT
    if sub:
        ws.merge_cells("A9:I9"); ws["A9"] = "  " + sub; ws["A9"].font = font(sz=10, color="555555")
    for col, w in zip("ABCDEFGHI", [8.4, 9.9, 22.6, 22, 8, 8, 12, 15.3, 30]):
        ws.column_dimensions[col].width = w

def table_head(ws, r):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for c, h in zip((1, 4, 5, 6, 7, 8, 9), ("구분", "규격", "단위", "수량", "단가", "금액", "비고")):
        x = ws.cell(r, c, h); x.font = font(True); x.alignment = CEN
    for c in range(1, 10): ws.cell(r, c).border = BOX; ws.cell(r, c).fill = GRAY

FOOT = ["* 전기, 배관·닥트 해체·복구 및 시운전은 별도입니다.",
        "* 1일 근무시간 8시간 기준이며, 오버타임 발생 시 추가 금액이 있습니다.",
        "* 크레인 사양·위치는 인양계획서와 구조검토 확정 후 조정될 수 있습니다."]

def detail(name, rows, work):
    ws = wb.create_sheet(name)
    header(ws, work)
    table_head(ws, 10)
    r = 11
    first = r
    labor = []
    groups = {}
    for g, item, spec, unit, q, p, note, is_labor in rows:
        groups.setdefault(g, []).append(r)
        ws.cell(r, 3, item); ws.cell(r, 4, spec); ws.cell(r, 5, unit)
        ws.cell(r, 6, q).font = font(color="0000FF")
        ws.cell(r, 7, p).font = font(color="0000FF")
        ws.cell(r, 8, f"=F{r}*G{r}")
        ws.cell(r, 9, note)
        if is_labor: labor.append(r)
        r += 1
    # 고용·산재보험 (인건비 × 3.3%)
    ins = r
    ws.cell(r, 3, "고용&산재보험"); ws.cell(r, 5, "%"); ws.cell(r, 6, INS * 100)
    ws.cell(r, 8, "=(" + "+".join(f"H{x}" for x in labor) + f")*{INS}"); ws.cell(r, 9, "인건비 × 3.3%")
    groups.setdefault("기타비용", []).append(r); r += 1
    # 기업이윤 = 목표금액 맞춤
    base_sum = sum(x[4] * x[5] for x in rows)
    lab_sum = sum(x[4] * x[5] for x in rows if x[7])
    profit = TARGETS[name] - base_sum - round(lab_sum * INS)
    ws.cell(r, 3, "기업이윤"); ws.cell(r, 5, "식"); ws.cell(r, 6, 1).font = font(color="0000FF")
    ws.cell(r, 7, profit).font = font(color="0000FF"); ws.cell(r, 8, f"=F{r}*G{r}")
    ws.cell(r, 9, "목표 금액 맞춤 조정 항목")
    groups.setdefault("기타비용", []).append(r); r += 1
    last = r - 1
    ws.cell(r, 3, "소    계"); ws.cell(r, 8, f"=SUM(H{first}:H{last})")
    for c in (3, 8): ws.cell(r, c).font = font(True)
    sub = r; r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(r, 1, "합                             계").font = font(True)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    ws.cell(r, 5, "(부가세 별도)").font = font(True); ws.cell(r, 5).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(r, 8, f"=H{sub}").font = font(True)
    tot = r
    # 구분 병합 (연속 행만)
    order = []
    for rr in range(first, last + 1):
        g = next(k for k, v in groups.items() if rr in v)
        if order and order[-1][0] == g and order[-1][2] == rr - 1: order[-1][2] = rr
        else: order.append([g, rr, rr])
    for g, a, b in order:
        ws.merge_cells(start_row=a, start_column=1, end_row=b, end_column=2)
        ws.cell(a, 1, g)
    for rr in range(first, tot + 1):
        for c in range(1, 10):
            x = ws.cell(rr, c); x.border = BOX
            if x.font.name != FN: x.font = font(x.font.b)
            x.alignment = LEFT if c == 9 else CEN
            if c in (7, 8): x.number_format = WON; x.alignment = Alignment(horizontal="right", vertical="center")
            if c == 6: x.number_format = '#,##0.0' if rr == ins else '#,##0'
        ws.row_dimensions[rr].height = 22
    for rr in (sub, tot):
        for c in range(1, 10): ws.cell(rr, c).fill = GRAY
    fr = tot + 2
    for t in FOOT:
        ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=9)
        ws.cell(fr, 1, t).font = font(sz=10); fr += 1
    ws.cell(fr + 1, 1, "파란 글씨 = 수량·단가 가정값(수정 가능). 기업이윤은 공종 목표 금액에 맞춰 넣은 값이라, 다른 단가를 바꾸면 소계가 달라져요.").font = font(sz=9, color="555555")
    ws.page_setup.paperSize = 9; ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 1
    ws.sheet_view.showGridLines = False
    return f"'{name}'!H{tot}" if " " in name else f"{name}!H{tot}"

WORK = "금강변압기 유한회사 마당 진공건조로 2분할체(35TON 개구부 측, 48TON) 회전·안착, 베이스 제작, 접합부 사상·용접"
header(cover, WORK, "공기: 총 11일 (베이스 3일 → 안착 1일 → 사상 2일 → 밀착 이동 2일 → 용접 3일)")
refs = {}
refs["도비"] = detail("도비", DOBI, "진공건조로 2기 회전·베이스 안착(1일) 및 용접면 밀착 이동(2일)")
refs["용접"] = detail("용접", WELD, "접합면 사상(2일) 및 접합부 용접(3일)")
refs["베이스"] = detail("베이스", BASE, "사다리형 일체 베이스 제작·설치, H-250×250 (4명 × 3일)")

table_head(cover, 11)
summary = [("1", "도비 작업", "회전·안착, 밀착 이동", "도비"),
           ("2", "용접 작업", "사상, 접합부 용접, 검사", "용접"),
           ("3", "베이스 작업", "사다리형 일체 베이스 제작·설치", "베이스")]
r = 12
for no, name, spec, key in summary:
    cover.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    cover.cell(r, 1, f"{no}. {name}"); cover.cell(r, 4, spec); cover.cell(r, 5, "식"); cover.cell(r, 6, 1)
    cover.cell(r, 7, f"={refs[key]}"); cover.cell(r, 8, f"=F{r}*G{r}"); cover.cell(r, 9, f"'{key}' 시트 참조")
    for c in (7,): cover.cell(r, c).font = font(color="008000")
    r += 1
first, last = 12, r - 1
cover.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
cover.cell(r, 1, "공 급 가 액"); cover.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
cover.cell(r, 5, "(부가세 별도)"); cover.cell(r, 8, f"=SUM(H{first}:H{last})"); sup = r; r += 1
cover.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
cover.cell(r, 1, "부 가 세 (10%)"); cover.cell(r, 8, f"=H{sup}*0.1"); vat = r; r += 1
cover.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
cover.cell(r, 1, "합                             계"); cover.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
cover.cell(r, 5, "(부가세 포함)"); cover.cell(r, 8, f"=H{sup}+H{vat}"); tot = r
for rr in range(12, tot + 1):
    for c in range(1, 10):
        x = cover.cell(rr, c); x.border = BOX
        b = rr >= sup and c in (1, 5, 8)
        x.font = font(b, color="008000" if (c == 7 and rr <= last) else "000000")
        x.alignment = LEFT if c == 9 else CEN
        if c in (7, 8): x.number_format = WON; x.alignment = Alignment(horizontal="right", vertical="center")
        if rr >= sup: x.fill = GRAY
    cover.row_dimensions[rr].height = 24
fr = tot + 2
for t in FOOT + ["* 베이스 H빔은 중고품 기준입니다. 신품 적용 시 3,997,657원이 추가됩니다.",
                 "* 접합부 용접 길이와 비파괴검사 범위는 발주처 확인 후 정산합니다."]:
    cover.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=9)
    cover.cell(fr, 1, t).font = font(sz=10); fr += 1
cover.page_setup.paperSize = 9; cover.sheet_properties.pageSetUpPr.fitToPage = True; cover.page_setup.fitToHeight = 1
cover.sheet_view.showGridLines = False
cover.column_dimensions["G"].width = 15

# ---- 베이스 자재 (H빔 구매 견적)
bs = wb.create_sheet("베이스자재")
bs["A1"] = "베이스 자재: H빔 H-250×250×9×14 (공급처 견적, 부가세 미포함)"; bs["A1"].font = font(True, 13)
hd = ["길이(m)", "수량(EA)", "총 길이(m)", "단위중량(kg/m)", "중량(kg)"]
for c, h in enumerate(hd, 1):
    x = bs.cell(3, c, h); x.font = font(True); x.fill = GRAY; x.border = BOX; x.alignment = CEN
for i, (ln, q) in enumerate([(10, 4), (5, 16), (0.6, 30)]):
    r = 4 + i
    bs.cell(r, 1, ln).font = font(color="0000FF"); bs.cell(r, 2, q).font = font(color="0000FF")
    bs.cell(r, 3, f"=A{r}*B{r}"); bs.cell(r, 4, 72.4).font = font(color="0000FF"); bs.cell(r, 5, f"=C{r}*D{r}")
    for c in range(1, 6): bs.cell(r, c).border = BOX; bs.cell(r, c).alignment = CEN
    bs.cell(r, 5).number_format = '#,##0'
bs.cell(7, 2, "합계").font = font(True); bs.cell(7, 3, "=SUM(C4:C6)"); bs.cell(7, 5, "=SUM(E4:E6)"); bs.cell(7, 5).number_format = '#,##0'
for c in range(1, 6): bs.cell(7, c).border = BOX; bs.cell(7, c).fill = GRAY; bs.cell(7, c).alignment = CEN
rows = [("신품", 13575440), ("중고 (견적 반영)", 9577783), ("운임", 160000), ("지게차 2일", 600000)]
bs.cell(9, 1, "구분").font = font(True); bs.cell(9, 2, "금액(원)").font = font(True); bs.cell(9, 3, "kg당(원)").font = font(True)
for i, (k, v) in enumerate(rows):
    r = 10 + i
    bs.cell(r, 1, k); bs.cell(r, 2, v).font = font(color="0000FF"); bs.cell(r, 2).number_format = WON
    if i < 2: bs.cell(r, 3, f"=B{r}/$E$7"); bs.cell(r, 3).number_format = '#,##0'
    for c in range(1, 4): bs.cell(r, c).border = BOX
bs.cell(14, 1, "신품 차액"); bs.cell(14, 2, "=B10-B11"); bs.cell(14, 2).number_format = WON
bs.cell(16, 1, "단위중량 72.4 kg/m는 KS H-250×250×9×14 공칭값이에요. 수량·금액은 사용자가 받은 공급처 견적이에요.").font = font(sz=9, color="555555")
for col, w in zip("ABCDE", [18, 16, 14, 16, 14]): bs.column_dimensions[col].width = w
bs.sheet_view.showGridLines = False

# ---- 공정표
ps = wb.create_sheet("공정표")
ps["A1"] = "공정표 (총 11일)"; ps["A1"].font = font(True, 14)
steps = [("베이스 제작·설치", "제작 기공 4명", 1, 3), ("회전·베이스 안착 (100t + 60t)", "도비 7명, 안전관리자", 4, 4),
         ("접합면 사상", "사상 기술자 2명, 조공 2명", 5, 6), ("용접면 밀착 이동", "도비 7명, 안전관리자", 7, 8),
         ("접합부 용접", "용접사 3명, 고급기술자 2명", 9, 11)]
ps.cell(3, 1, "공정").font = font(True); ps.cell(3, 2, "투입 인원").font = font(True)
for d in range(1, 12):
    x = ps.cell(3, 2 + d, f"{d}일"); x.font = font(True); x.alignment = CEN
    ps.column_dimensions[ps.cell(3, 2 + d).column_letter].width = 6
BAR = PatternFill("solid", fgColor="E8742A")
for i, (nm, crew, a, b) in enumerate(steps):
    rr = 4 + i
    ps.cell(rr, 1, nm).font = font(); ps.cell(rr, 2, crew).font = font()
    for d in range(1, 12):
        x = ps.cell(rr, 2 + d); x.border = BOX
        if a <= d <= b: x.fill = BAR
    ps.cell(rr, 1).border = BOX; ps.cell(rr, 2).border = BOX
for c in range(1, 14): ps.cell(3, c).border = BOX; ps.cell(3, c).fill = GRAY
ps.column_dimensions["A"].width = 30; ps.column_dimensions["B"].width = 28
ps.cell(10, 1, "사용자 제시 공정 기준. 베이스 제작은 공장 선행 제작 시 공기를 줄일 수 있어요.").font = font(sz=9, color="555555")
ps.sheet_view.showGridLines = False

wb.save(OUT)
print(OUT)
